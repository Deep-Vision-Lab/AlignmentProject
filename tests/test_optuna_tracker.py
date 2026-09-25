import json

from openpyxl import load_workbook

from optuna_tracker import initialize, save_trial


def _config():
    return dict(fusion_mode='concat', use_gated_fusion=0, local_dropout=.1,
                embedding_dim=128, sigreg_weight=0., dtw_gamma=.05,
                transformer_layers=5, transformer_heads=1,
                window_width=32, window_stride=16, seed=42)


def _history(losses, f1=None):
    return [dict(epoch=index, train=dict(total=loss + .5, positive_dtw=loss + .5,
                   seconds=2., gradients=dict(cnn_mean=.1, transformer_mean=.2,
                   fusion_mean=.3)), validation=dict(total=loss, positive_dtw=loss,
                   alignment_f1=f1, seconds=1.), model_parameter_count=123)
            for index, loss in enumerate(losses, 1)]


def test_tracker_saves_every_epoch_and_physically_sorts(tmp_path):
    path = tmp_path / 'tracker.xlsx'
    initialize(path, {'fusion': ['concat', 'sum']})
    save_trial(path, trial_id=0, status='COMPLETE', config=_config(),
               history=_history([2., 1.]), search_space={'fusion': ['concat', 'sum']})
    save_trial(path, trial_id=1, status='COMPLETE', config=_config(),
               history=_history([3., .8]), search_space={'fusion': ['concat', 'sum']})
    save_trial(path, trial_id=2, status='PRUNED', config=_config(),
               history=_history([.1]), search_space={'fusion': ['concat', 'sum']})
    book = load_workbook(path, data_only=True)
    assert book.sheetnames == ['RankedResults', 'Trials', 'EpochMetrics', 'Dashboard', 'SearchSpace']
    rows = list(book['RankedResults'].iter_rows(values_only=True))
    head = {value: index for index, value in enumerate(rows[0])}
    assert [row[head['Trial ID']] for row in rows[1:]] == [1, 0, 2]
    assert [row[head['Rank']] for row in rows[1:]] == [1, 2, None]
    assert rows[1][head['Learning_Improvement']] == 2.2
    epoch_rows = list(book['EpochMetrics'].iter_rows(values_only=True))
    assert len(epoch_rows) == 11  # Header + (2 + 2 + 1) epochs x 2 splits.
    assert {(row[0], row[1], row[2]) for row in epoch_rows[1:]} == {
        (0, 1, 'train'), (0, 1, 'val'), (0, 2, 'train'), (0, 2, 'val'),
        (1, 1, 'train'), (1, 1, 'val'), (1, 2, 'train'), (1, 2, 'val'),
        (2, 1, 'train'), (2, 1, 'val')}
    assert book['Dashboard']['B3'].value == 1
    assert book['Dashboard']['B7'].value == 2
    assert book['Dashboard']['B8'].value == 1
    assert book['SearchSpace']['B2'].value == json.dumps(['concat', 'sum'])
    book.close()


def test_f1_priority_and_improvement_tie_break(tmp_path):
    path = tmp_path / 'tracker.xlsx'
    for trial_id, losses, f1 in ((0, [2., .1], None), (1, [2., 1.], .6),
                                 (2, [2., 1.], .6)):
        history = _history(losses, f1)
        if trial_id == 1:
            history[0]['validation']['alignment_f1'] = .3
        if trial_id == 2:
            history[0]['validation']['alignment_f1'] = .5
        save_trial(path, trial_id=trial_id, status='COMPLETE', config=_config(),
                   history=history)
    book = load_workbook(path, data_only=True)
    assert [row[1] for row in book['RankedResults'].iter_rows(min_row=2, values_only=True)] == [1, 2, 0]
    book.close()
