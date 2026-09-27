"""Small Optuna/SQLite/Excel integration test without model training."""
import io
import json
import os
from contextlib import redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
import optuna

from scripts.train import run_optuna


class _ChoiceTrial:
    def __init__(self, dimension):
        self.dimension = dimension
        self.seen = {}

    def suggest_categorical(self, name, choices):
        self.seen[name] = tuple(choices)
        return self.dimension if name == 'vector_dimension' else choices[0]


class _DepthTrial:
    """Fake trial that forces a specific cnn_layers choice."""

    def __init__(self, depth):
        self.depth = depth
        self.seen = {}

    def suggest_categorical(self, name, choices):
        self.seen[name] = tuple(choices)
        return self.depth if name == 'cnn_layers' else choices[0]


class OptunaSmokeTest(unittest.TestCase):
    def test_conditional_heads_for_every_dimension(self):
        for dimension in (64, 128, 256):
            with self.subTest(dimension=dimension):
                trial = _ChoiceTrial(dimension)
                config = run_optuna._suggest_config(
                    trial, {}, run_optuna.DEFAULT_SEARCH_SPACE)
                self.assertEqual(trial.seen['transformer_heads'], (1, 2, 4))
                self.assertNotIn(3, trial.seen['transformer_heads'])
                self.assertEqual(config['embedding_dim'], dimension)
                self.assertEqual(dimension % config['transformer_heads'], 0)
                self.assertTrue(all(dimension % head == 0
                                    for head in trial.seen['transformer_heads']))
                self.assertEqual(set(trial.seen), set(run_optuna.DEFAULT_SEARCH_SPACE))

    def test_cnn_layers_sampled_only_for_simple_cnn(self):
        trial = _DepthTrial(4)
        config = run_optuna._suggest_config(trial, {'cnn_type': 'simple'},
                                            run_optuna.DEFAULT_SEARCH_SPACE)
        self.assertEqual(trial.seen['cnn_layers'], (1, 2, 3, 4, 5))
        self.assertEqual(config['cnn_layers'], 4)

        resnet_trial = _DepthTrial(4)
        config = run_optuna._suggest_config(resnet_trial,
                                            {'cnn_type': 'resnet18', 'cnn_layers': 3},
                                            run_optuna.DEFAULT_SEARCH_SPACE)
        self.assertNotIn('cnn_layers', resnet_trial.seen)
        self.assertEqual(config['cnn_layers'], 3)  # fixed; ResNet18 depth never changes

    def test_every_simple_cnn_depth_builds_and_trains(self):
        from dataclasses import replace

        import torch

        from parameters import Config
        from train import build_model

        for depth in (1, 2, 3, 4, 5):
            with self.subTest(depth=depth):
                trial = _DepthTrial(depth)
                suggested = run_optuna._suggest_config(trial, {'cnn_type': 'simple'},
                                                       run_optuna.DEFAULT_SEARCH_SPACE)
                self.assertEqual(suggested['cnn_layers'], depth)  # sampled -> Config path
                config = replace(Config(**suggested), cnn_pretrained=False,
                                 image_height=32, image_width=64)
                self.assertEqual(config.cnn_layers, depth)
                model = build_model(config)
                convs = [module for module in model.cnn.backbone
                         if isinstance(module, torch.nn.Conv2d)]
                self.assertEqual(len(convs), depth)  # correct CNN depth built
                output = model(torch.randn(1, 1, 32, 64))
                self.assertTrue(torch.isfinite(output['fused']).all())
                loss = (output['fused'] * torch.randn_like(output['fused'])).sum()
                self.assertTrue(torch.isfinite(loss))
                loss.backward()
                self.assertGreater(model.cnn.projection[0].weight.grad.abs().sum().item(), 0)
                for conv in convs:
                    self.assertGreater(conv.weight.grad.abs().sum().item(), 0)

    def test_resnet18_depth_stays_fixed(self):
        from dataclasses import replace

        import torch

        from parameters import Config
        from train import build_model

        trial = _DepthTrial(5)
        suggested = run_optuna._suggest_config(trial,
                                               {'cnn_type': 'resnet18', 'cnn_layers': 3},
                                               run_optuna.DEFAULT_SEARCH_SPACE)
        self.assertNotIn('cnn_layers', trial.seen)
        config = replace(Config(**suggested), cnn_pretrained=False,
                         image_height=128, image_width=64)
        model = build_model(config)
        self.assertTrue(hasattr(model.cnn.backbone, 'layer4'))  # full ResNet18 kept
        self.assertIsNone(model.cnn.num_layers)
        output = model(torch.randn(1, 1, 128, 64))
        self.assertTrue(torch.isfinite(output['fused']).all())

    def test_full_epoch_failure_excel_and_sqlite(self):
        optuna.logging.set_verbosity(optuna.logging.WARNING)
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            tracker = root / 'tracker.xlsx'
            database = root / 'study.db'
            checkpoints = root / 'checkpoints'
            logs = root / 'logs'
            study_name = 'no_prune_smoke'
            trained_epochs = {}

            def fake_training(argv):
                def argument(name):
                    return argv[argv.index(name) + 1]

                run_name = argument('--run-name')
                configured_epochs = int(argument('--epochs'))
                output = Path(argument('--output-root')) / run_name
                output.mkdir(parents=True)
                if run_name == 'trial_00001':
                    self.assertTrue(tracker.exists())
                    prior = load_workbook(tracker, read_only=True, data_only=True)
                    self.assertEqual(prior['Dashboard']['B7'].value, 1)
                    prior.close()
                history = []
                for epoch in range(1, configured_epochs + 1):
                    history.append(dict(epoch=epoch, model_parameter_count=100,
                        train=dict(total=2.0 / epoch, positive_dtw=2.0 / epoch,
                                   seconds=.1, gradients={}),
                        validation=dict(total=1.0 / epoch,
                                        positive_dtw=1.0 / epoch, seconds=.1)))
                    (output / 'history.json').write_text(json.dumps(history))
                    trained_epochs[run_name] = epoch
                    if run_name == 'trial_00001':
                        raise RuntimeError('simulated real training error')
                (output / 'checkpoint_latest.pt').write_bytes(b'smoke checkpoint')

            arguments = ['--dataset', str(root), '--n-trials', '2', '--epochs', '3',
                         '--tracker', str(tracker), '--checkpoint-root', str(checkpoints),
                         '--log-root', str(logs), '--storage', f'sqlite:///{database}',
                         '--study-name', study_name]
            with patch.object(run_optuna, 'train_main', fake_training):
                run_optuna.main(arguments)

            study = optuna.load_study(study_name=study_name,
                                      storage=f'sqlite:///{database}')
            self.assertEqual([trial.state.name for trial in study.trials],
                             ['COMPLETE', 'FAIL'])
            self.assertEqual(trained_epochs, {'trial_00000': 3, 'trial_00001': 1})
            for trial in study.trials:
                dimension = trial.params['vector_dimension']
                heads = trial.params['transformer_heads']
                self.assertEqual(dimension % heads, 0)
                self.assertNotEqual(heads, 3)
                self.assertEqual(tuple(trial.distributions['transformer_heads'].choices),
                                 (1, 2, 4))
                # Default base config is resnet18: CNN depth stays fixed and is
                # never sampled as a study parameter.
                self.assertNotIn('cnn_layers', trial.params)
                self.assertFalse(trial.intermediate_values)
                log_text = (logs / f'trial_{trial.number:05d}.log').read_text()
                self.assertIn('CNN type: resnet18', log_text)
                self.assertIn('CNN layers: fixed', log_text)
            self.assertIn('simulated real training error',
                          study.trials[1].user_attrs['error'])

            workbook = load_workbook(tracker, data_only=True)
            self.assertEqual(workbook.sheetnames,
                ['RankedResults', 'Trials', 'EpochMetrics', 'Dashboard', 'SearchSpace'])
            ranked = list(workbook['RankedResults'].iter_rows(values_only=True))
            self.assertEqual(ranked[1][0:3], (1, f'{study_name}:0', 'COMPLETE'))
            self.assertEqual(ranked[2][0:3], (None, f'{study_name}:1', 'FAILED'))
            epochs = list(workbook['EpochMetrics'].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(epochs), 8)
            self.assertEqual({(row[0], row[1], row[2]) for row in epochs},
                {(f'{study_name}:0', epoch, split)
                 for epoch in (1, 2, 3) for split in ('train', 'val')} |
                {(f'{study_name}:1', 1, split) for split in ('train', 'val')})
            self.assertEqual(workbook['Dashboard']['B7'].value, 1)
            self.assertEqual(workbook['Dashboard']['B8'].value, 0)
            self.assertEqual(workbook['Dashboard']['B9'].value, 1)
            workbook.close()
            self.assertTrue(database.exists())

    def test_simple_cnn_study_searches_depth_and_records_it(self):
        optuna.logging.set_verbosity(optuna.logging.WARNING)
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            tracker = root / 'tracker.xlsx'
            database = root / 'study.db'
            checkpoints = root / 'checkpoints'
            logs = root / 'logs'
            study_name = 'simple_cnn_depth_smoke'
            base_config = root / 'base.json'
            base_config.write_text(json.dumps({'cnn_type': 'simple',
                                               'cnn_pretrained': False}))
            seen_argv = {}

            def fake_training(argv):
                def argument(name):
                    return argv[argv.index(name) + 1]

                run_name = argument('--run-name')
                seen_argv[run_name] = argv
                output = Path(argument('--output-root')) / run_name
                output.mkdir(parents=True)
                history = [dict(epoch=epoch, model_parameter_count=100,
                                train=dict(total=2.0, positive_dtw=2.0, seconds=.1,
                                           gradients={}),
                                validation=dict(total=1.0, positive_dtw=1.0, seconds=.1))
                           for epoch in (1, 2)]
                (output / 'history.json').write_text(json.dumps(history))

            arguments = ['--dataset', str(root), '--n-trials', '3', '--epochs', '2',
                         '--base-config', str(base_config),
                         '--tracker', str(tracker), '--checkpoint-root', str(checkpoints),
                         '--log-root', str(logs), '--storage', f'sqlite:///{database}',
                         '--study-name', study_name]
            with patch.object(run_optuna, 'train_main', fake_training):
                run_optuna.main(arguments)

            study = optuna.load_study(study_name=study_name,
                                      storage=f'sqlite:///{database}')
            self.assertEqual([trial.state.name for trial in study.trials],
                             ['COMPLETE'] * 3)
            workbook = load_workbook(tracker, data_only=True)
            trial_rows = list(workbook['Trials'].iter_rows(values_only=True))
            headers = {value: index for index, value in enumerate(trial_rows[0])}
            self.assertIn('CNN type', headers)
            self.assertIn('CNN layers', headers)
            for trial in study.trials:
                depth = trial.params['cnn_layers']  # sampled for the simple CNN
                self.assertIn(depth, (1, 2, 3, 4, 5))
                self.assertEqual(tuple(trial.distributions['cnn_layers'].choices),
                                 (1, 2, 3, 4, 5))
                self.assertEqual(trial.params['vector_dimension']
                                 % trial.params['transformer_heads'], 0)
                self.assertFalse(trial.intermediate_values)  # no pruning
                run_name = f'trial_{trial.number:05d}'
                argv = seen_argv[run_name]  # sampled value reaches the training CLI
                self.assertEqual(int(argv[argv.index('--cnn-layers') + 1]), depth)
                self.assertEqual(argv[argv.index('--cnn-type') + 1], 'simple')
                config = trial.user_attrs['config']
                self.assertEqual(config['cnn_layers'], depth)  # and reaches Config
                log_text = (logs / f'{run_name}.log').read_text()
                self.assertIn('CNN type: simple', log_text)
                self.assertIn(f'CNN layers: {depth}', log_text)
                row = next(row for row in trial_rows[1:]
                           if row[headers['Trial ID']] == f'{study_name}:{trial.number}')
                self.assertEqual(row[headers['CNN type']], 'simple')
                self.assertEqual(row[headers['CNN layers']], depth)
            space = {row[0]: row[1] for row in
                     workbook['SearchSpace'].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(json.loads(space['cnn_layers']), [1, 2, 3, 4, 5])
            workbook.close()

    def test_study_without_cnn_layers_distribution_is_rejected(self):
        optuna.logging.set_verbosity(optuna.logging.WARNING)
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            database = root / 'legacy_layers.db'
            storage = f'sqlite:///{database}'
            legacy = optuna.create_study(study_name='legacy_layers', storage=storage,
                                         direction='maximize')
            legacy.optimize(lambda trial: float(
                trial.suggest_categorical('vector_dimension', [64, 128, 256]) /
                trial.suggest_categorical('transformer_heads', [1, 2, 4])), n_trials=1)
            base_config = root / 'base.json'
            base_config.write_text(json.dumps({'cnn_type': 'simple'}))
            with self.assertRaisesRegex(RuntimeError, 'new --study-name and --storage'):
                run_optuna.main(['--dataset', str(root), '--n-trials', '1',
                    '--base-config', str(base_config),
                    '--tracker', str(root / 'tracker.xlsx'), '--storage', storage,
                    '--study-name', 'legacy_layers', '--checkpoint-root', str(root / 'weights'),
                    '--log-root', str(root / 'logs')])

    def test_job_name_sanitization(self):
        self.assertEqual(run_optuna._sanitize_job_name('cnn_layers_optuna'),
                         'cnn_layers_optuna')
        self.assertEqual(run_optuna._sanitize_job_name('my job/x:y'), 'my_job_x_y')
        self.assertEqual(run_optuna._sanitize_job_name('a*b?c'), 'a_b_c')
        self.assertEqual(run_optuna._sanitize_job_name('..'), 'optuna')
        self.assertEqual(run_optuna._sanitize_job_name(''), 'optuna')

    def _run_fake_job(self, working_dir, extra_args, env, n_trials=2):
        """Run one fake Optuna job with working_dir as the current directory."""
        def fake_training(argv):
            def argument(name):
                return argv[argv.index(name) + 1]

            output = Path(argument('--output-root')) / argument('--run-name')
            output.mkdir(parents=True)
            history = [dict(epoch=1, model_parameter_count=100,
                            train=dict(total=2.0, positive_dtw=2.0, seconds=.1,
                                       gradients={}),
                            validation=dict(total=1.0, positive_dtw=1.0, seconds=.1))]
            (output / 'history.json').write_text(json.dumps(history))

        previous = os.getcwd()
        os.chdir(working_dir)
        try:
            with patch.dict(os.environ, clear=False) as patched:
                patched.pop('SLURM_JOB_NAME', None)
                patched.pop('SLURM_JOB_ID', None)
                patched.update(env)
                with patch.object(run_optuna, 'train_main', fake_training):
                    stdout = io.StringIO()
                    with redirect_stdout(stdout):
                        run_optuna.main(['--dataset', '.', '--n-trials', str(n_trials),
                                         '--epochs', '1'] + extra_args)
            return stdout.getvalue()
        finally:
            os.chdir(previous)

    def test_slurm_job_name_organizes_all_outputs(self):
        optuna.logging.set_verbosity(optuna.logging.WARNING)
        with TemporaryDirectory() as temporary:
            header = self._run_fake_job(
                Path(temporary), ['--study-name', 'alignment_optuna'],
                {'SLURM_JOB_NAME': 'test_optuna_job', 'SLURM_JOB_ID': '12345'})
            root = Path(temporary) / 'results/test_optuna_job'
            self.assertIn('SLURM job name : test_optuna_job', header)
            self.assertIn('SLURM job ID   : 12345', header)
            self.assertIn('Study name     : alignment_optuna', header)
            self.assertTrue((root / 'optuna_alignment_experiment_tracker.xlsx').is_file())
            self.assertTrue((root / 'study.db').is_file())
            # Every trial gets its own folder; no trial overwrites another.
            self.assertTrue((root / 'checkpoints/trial_00000/history.json').is_file())
            self.assertTrue((root / 'checkpoints/trial_00001/history.json').is_file())
            self.assertTrue((root / 'logs/trial_00000.log').is_file())
            self.assertTrue((root / 'logs/trial_00001.log').is_file())
            space = json.loads((root / 'search_space.json').read_text())
            self.assertEqual(space['study_name'], 'alignment_optuna')
            self.assertEqual(space['job_name'], 'test_optuna_job')
            self.assertEqual(space['slurm_job_id'], '12345')
            self.assertEqual(space['search_space']['cnn_layers'], [1, 2, 3, 4, 5])
            self.assertEqual(set(space['search_space']),
                             set(run_optuna.DEFAULT_SEARCH_SPACE))
            workbook = load_workbook(root / 'optuna_alignment_experiment_tracker.xlsx',
                                     read_only=True)
            self.assertIn('RankedResults', workbook.sheetnames)
            workbook.close()

    def test_explicit_and_fallback_job_names_create_separate_roots(self):
        optuna.logging.set_verbosity(optuna.logging.WARNING)
        with TemporaryDirectory() as temporary:
            temporary = Path(temporary)
            header = self._run_fake_job(
                temporary, ['--study-name', 'fallback_study', '--job-name', 'local_test'],
                {})
            explicit = temporary / 'results/local_test'
            self.assertTrue((explicit / 'study.db').is_file())
            self.assertTrue((explicit / 'checkpoints/trial_00000/history.json').is_file())
            self.assertIn('SLURM job name : N/A', header)
            self.assertIn('SLURM job ID   : N/A', header)
            # Fallback: no SLURM and no --job-name -> study name, separate folder.
            self._run_fake_job(temporary, ['--study-name', 'fallback_study'], {})
            fallback = temporary / 'results/fallback_study'
            self.assertTrue((fallback / 'study.db').is_file())
            self.assertTrue((fallback / 'logs/trial_00000.log').is_file())
            # Different job names create completely separate result folders.
            self.assertTrue(explicit.is_dir() and fallback.is_dir())
            self.assertNotEqual(explicit, fallback)

    def test_explicit_path_overrides_bypass_the_job_folder(self):
        optuna.logging.set_verbosity(optuna.logging.WARNING)
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            self._run_fake_job(root, ['--study-name', 'override_study',
                                      '--tracker', str(root / 'custom/tracker.xlsx'),
                                      '--checkpoint-root', str(root / 'custom/checkpoints'),
                                      '--log-root', str(root / 'custom/logs'),
                                      '--storage', f'sqlite:///{root / "custom/study.db"}'],
                               {}, n_trials=1)
            self.assertTrue((root / 'custom/tracker.xlsx').is_file())
            self.assertTrue((root / 'custom/study.db').is_file())
            self.assertTrue((root / 'custom/checkpoints/trial_00000/history.json')
                            .is_file())
            self.assertTrue((root / 'custom/logs/trial_00000.log').is_file())
            # search_space.json follows the tracker; nothing writes into results/.
            self.assertTrue((root / 'custom/search_space.json').is_file())
            self.assertFalse((root / 'results').exists())

    def test_legacy_head_distribution_is_rejected(self):
        optuna.logging.set_verbosity(optuna.logging.WARNING)
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            database = root / 'legacy.db'
            storage = f'sqlite:///{database}'
            legacy = optuna.create_study(study_name='alignment', storage=storage,
                                         direction='maximize')
            legacy.optimize(lambda trial: float(trial.suggest_categorical(
                'transformer_heads', [1, 2, 3])), n_trials=1)
            with self.assertRaisesRegex(RuntimeError, 'new --study-name and --storage'):
                run_optuna.main(['--dataset', str(root), '--n-trials', '1',
                    '--tracker', str(root / 'tracker.xlsx'), '--storage', storage,
                    '--study-name', 'alignment', '--checkpoint-root', str(root / 'weights'),
                    '--log-root', str(root / 'logs')])


if __name__ == '__main__':
    unittest.main()
