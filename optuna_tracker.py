"""Crash-safe Excel snapshot of completed Optuna trial state.

The workbook is rebuilt under a file lock and atomically replaced after each
trial. The trial history supplied here is the saved train.py history.json.
"""
from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime, timezone
from functools import cmp_to_key
import json
import math
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference


DEFAULT_TRACKER = Path('results/optuna/optuna_alignment_experiment_tracker.xlsx')
METRICS = (
    'Total loss', 'DTW loss', 'SIGReg loss', 'Alignment F1', 'Mask IoU',
    'Positive similarity', 'Negative similarity', 'Similarity margin',
    'Local gradient norm', 'Contextual gradient norm', 'Fusion/gate gradient norm',
    'Gate mean', 'Gate std', 'Gate min', 'Gate max', 'Learning rate',
    'Local vector norm', 'Contextual vector norm', 'Fused vector norm',
    'Similarity matrix min', 'Similarity matrix max', 'Similarity matrix mean',
    'Similarity matrix std', 'Epoch time (s)', 'GPU memory (MB)',
)
METRIC_KEYS = (
    'total', 'positive_dtw', 'sigreg', 'alignment_f1', 'mask_iou',
    'positive_similarity', 'negative_similarity', 'similarity_margin',
    'local_gradient_norm', 'contextual_gradient_norm', 'fusion_gradient_norm',
    'gate_mean', 'gate_std', 'gate_min', 'gate_max', 'learning_rate',
    'local_vector_norm', 'contextual_vector_norm', 'fused_vector_norm',
    'similarity_min', 'similarity_max', 'similarity_mean', 'similarity_std',
    'seconds', 'gpu_memory_mb',
)
CONFIG_HEADERS = (
    'Trial ID', 'Status', 'Fusion', 'Dropout', 'Vector dimension', 'SIGReg ON/OFF',
    'DTW gamma', 'Transformer layers', 'Transformer heads', 'Window size',
    'Stride ratio', 'Actual stride pixels',
)
SUMMARY_HEADERS = (
    'First objective', 'Final objective', 'Best objective', 'Best epoch',
    'Epoch count', 'Learning_Improvement', 'Objective metric',
    'Model parameter count', 'Mean epoch time (s)', 'Peak GPU memory (MB)',
    'Checkpoint path', 'Log path', 'Seed',
)
RANKED_HEADERS = ('Rank',) + CONFIG_HEADERS + tuple(
    f'{stage} {split} — {metric}'
    for stage in ('Epoch 1', 'Final epoch') for split in ('TRAIN', 'VALIDATION')
    for metric in METRICS
) + SUMMARY_HEADERS
TRIAL_HEADERS = CONFIG_HEADERS + SUMMARY_HEADERS + ('Config JSON', 'Error', 'Updated UTC')
EPOCH_HEADERS = ('Trial_ID', 'Epoch', 'Split') + METRICS


def _finite(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def metric_values(stats, split):
    stats = stats or {}
    gradients = stats.get('gradients') or {} if split == 'train' else {}
    values = dict(stats)
    values.update(local_gradient_norm=gradients.get('cnn_mean'),
                  contextual_gradient_norm=gradients.get('transformer_mean'),
                  fusion_gradient_norm=gradients.get('fusion_mean'))
    return tuple(_finite(values.get(key)) for key in METRIC_KEYS)


def objective(stats):
    """Return the raw final validation metric and its source."""
    if not stats:
        return None, None
    for key, label in (('alignment_f1', 'Alignment F1'), ('mask_iou', 'Mask IoU')):
        value = _finite(stats.get(key))
        if value is not None:
            return value, label
    loss = _finite(stats.get('total'))
    return (loss, 'Total loss') if loss is not None else (None, None)


def ranking_score(stats):
    """Higher is always better, for Optuna/pruning and improvement arithmetic."""
    value, metric = objective(stats)
    return (-value if metric == 'Total loss' else value, metric) if value is not None else (None, None)


def _config_values(trial_id, status, config):
    config = config or {}
    stride = config.get('window_stride')
    width = config.get('window_width')
    fusion = 'gated_sum' if config.get('use_gated_fusion') else config.get('fusion_mode')
    return (trial_id, status, fusion, config.get('local_dropout'),
            config.get('embedding_dim'), 'ON' if config.get('sigreg_weight') else 'OFF',
            config.get('dtw_gamma'), config.get('transformer_layers'),
            config.get('transformer_heads'), width,
            stride / width if stride is not None and width else None, stride)


def _summary(history, checkpoint, log_path, seed):
    first = history[0]['validation'] if history else None
    last = history[-1]['validation'] if history else None
    first_score, first_metric = objective(first)
    final_score, final_metric = objective(last)
    scored = [(ranking_score(item.get('validation'))[0], item['epoch'])
              for item in history if objective(item.get('validation'))[1] == final_metric]
    scored = [(score, epoch) for score, epoch in scored if score is not None]
    best_score, best_epoch = (max(scored, key=lambda pair: (pair[0], -pair[1]))
                              if scored else (None, None))
    if best_score is not None and final_metric == 'Total loss':
        best_score = -best_score
    times = [(item.get('train') or {}).get('seconds', 0) +
             (item.get('validation') or {}).get('seconds', 0) for item in history]
    memories = [_finite((item.get(split) or {}).get('gpu_memory_mb'))
                for item in history for split in ('train', 'validation')]
    memories = [value for value in memories if value is not None]
    improvement = ((first_score - final_score if final_metric == 'Total loss'
                    else final_score - first_score)
                   if final_score is not None and first_score is not None and
                   final_metric == first_metric else None)
    return (first_score, final_score, best_score, best_epoch, len(history), improvement,
            final_metric, history[-1].get('model_parameter_count') if history else None,
            sum(times) / len(times) if times else None, max(memories) if memories else None,
            str(checkpoint) if checkpoint and Path(checkpoint).exists() else None,
            str(log_path) if log_path else None, seed)


def _ranked_row(trial):
    history = trial['history']
    first = history[0] if history else {}
    last = history[-1] if history else {}
    stages = (first, last)
    metrics = tuple(value for entry in stages for split in ('train', 'validation')
                    for value in metric_values(entry.get(split), split))
    return (_config_values(trial['id'], trial['status'], trial['config']) + metrics +
            _summary(history, trial.get('checkpoint'), trial.get('log_path'),
                     trial['config'].get('seed')))


def _sort_trials(left, right):
    status_order = {'COMPLETE': 0, 'PRUNED': 1, 'FAILED': 2}
    a, b = status_order.get(left['status'], 3), status_order.get(right['status'], 3)
    if a != b:
        return -1 if a < b else 1
    if a == 0:
        la = _summary(left['history'], None, None, None)
        lb = _summary(right['history'], None, None, None)
        metric_priority = {'Alignment F1': 0, 'Mask IoU': 1, 'Total loss': 2}
        pa = metric_priority.get(la[6], 3)
        pb = metric_priority.get(lb[6], 3)
        if pa != pb:
            return -1 if pa < pb else 1
        sa, sb = la[1], lb[1]
        if sa is None or sb is None:
            return -1 if sa is not None else (1 if sb is not None else 0)
        if abs(sa - sb) > 1e-4:
            return -1 if (sa < sb if la[6] == 'Total loss' else sa > sb) else 1
        ia, ib = la[5] or 0, lb[5] or 0
        if ia != ib:
            return -1 if ia > ib else 1
        if sa != sb:
            return -1 if (sa < sb if la[6] == 'Total loss' else sa > sb) else 1
    return (left['id'] > right['id']) - (left['id'] < right['id'])


def _style_table(sheet, headers, rows, color='16324F'):
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'D2' if sheet.title == 'RankedResults' else 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 54 if sheet.title == 'RankedResults' else 38
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor=color)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for column, header in enumerate(headers, 1):
        width = 17 if len(header) > 24 else max(13, min(26, len(header) + 3))
        if 'path' in header.lower():
            width = 42
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='EDF3F8')
            if isinstance(cell.value, float):
                cell.number_format = '0.0000'


def _dashboard(sheet, ranked):
    sheet.sheet_view.showGridLines = False
    sheet['A1'] = 'OPTUNA ALIGNMENT · VALIDATION LEADERBOARD'
    sheet['A1'].font = Font(size=18, bold=True, color='16324F')
    complete = [trial for trial in ranked if trial['status'] == 'COMPLETE']
    counts = {status: sum(trial['status'] == status for trial in ranked)
              for status in ('COMPLETE', 'PRUNED', 'FAILED')}
    best = complete[0] if complete else None
    score = _summary(best['history'], None, None, None) if best else None
    best_config = None
    if best:
        cfg = _config_values(best['id'], best['status'], best['config'])
        best_config = (f'Fusion={cfg[2]}; Dropout={cfg[3]}; Dim={cfg[4]}; '
                       f'SIGReg={cfg[5]}; gamma={cfg[6]}; Layers={cfg[7]}; '
                       f'Heads={cfg[8]}; Window={cfg[9]}; Stride={cfg[11]} px')
    summary = [('Best trial', best['id'] if best else 'No completed trials'),
               ('Best configuration', best_config),
               ('Best final validation score', score[1] if score else None),
               ('Score metric', score[6] if score else None),
               ('Completed', counts['COMPLETE']), ('Pruned', counts['PRUNED']),
               ('Failed', counts['FAILED'])]
    for row, (label, value) in enumerate(summary, 3):
        sheet.cell(row, 1, label).font = Font(bold=True, color='16324F')
        sheet.cell(row, 2, value)
    sheet['A12'] = 'TOP 10 CONFIGURATIONS · FINAL VALIDATION'
    sheet['A12'].font = Font(bold=True, color='16324F', size=13)
    headers = ('Rank', 'Trial ID', 'Fusion', 'Dropout', 'Vector dimension',
               'SIGReg', 'DTW gamma', 'Layers', 'Heads', 'Window', 'Stride',
               'Score', 'Metric', 'Improvement')
    for col, label in enumerate(headers, 1):
        cell = sheet.cell(13, col, label)
        cell.fill = PatternFill('solid', fgColor='16324F')
        cell.font = Font(bold=True, color='FFFFFF')
    for rank, trial in enumerate(complete[:10], 1):
        cfg = _config_values(trial['id'], trial['status'], trial['config'])
        summary_values = _summary(trial['history'], None, None, None)
        row = (rank, cfg[0], cfg[2], cfg[3], cfg[4], cfg[5], cfg[6], cfg[7],
               cfg[8], cfg[9], cfg[11], summary_values[1], summary_values[6],
               summary_values[5])
        for col, value in enumerate(row, 1):
            sheet.cell(rank + 13, col, value)
    sheet.column_dimensions['A'].width = 29
    sheet.column_dimensions['B'].width = 44
    for col in range(3, 15):
        sheet.column_dimensions[get_column_letter(col)].width = 17
    sheet['B4'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[4].height = 72
    if complete:
        chart = BarChart()
        chart.title = 'Top 10 validation score'
        chart.y_axis.title = ('Lower is better' if score and score[6] == 'Total loss'
                              else 'Higher is better')
        chart.add_data(Reference(sheet, min_col=12, min_row=13,
                                 max_row=13 + min(10, len(complete))), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=2, min_row=14,
                                       max_row=13 + min(10, len(complete))))
        chart.width, chart.height = 18, 8
        sheet.add_chart(chart, 'A26')


@contextmanager
def _locked(path):
    import fcntl
    path.parent.mkdir(parents=True, exist_ok=True)
    with (path.parent / (path.name + '.lock')).open('a+b') as lock:
        fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
        try:
            yield
        finally:
            fcntl.flock(lock.fileno(), fcntl.LOCK_UN)


def _read_existing(path):
    if not path.exists():
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    trials = {}
    if 'Trials' in workbook and 'EpochMetrics' in workbook:
        trial_rows = workbook['Trials'].iter_rows(values_only=True)
        trial_headers = next(trial_rows, ())
        for row in trial_rows:
            data = dict(zip(trial_headers, row))
            if data.get('Trial ID') is None:
                continue
            trials[data['Trial ID']] = dict(id=data['Trial ID'], status=data['Status'],
                config=json.loads(data['Config JSON']) if data.get('Config JSON') else {},
                history=[], checkpoint=data.get('Checkpoint path'),
                log_path=data.get('Log path'), error=data.get('Error'))
            cfg = trials[data['Trial ID']]['config']
            if not cfg:
                cfg.update(fusion_mode=data.get('Fusion'),
                           use_gated_fusion=int(data.get('Fusion') == 'gated_sum'),
                           local_dropout=data.get('Dropout'),
                           embedding_dim=data.get('Vector dimension'),
                           sigreg_weight=1.0 if data.get('SIGReg ON/OFF') == 'ON' else 0.0,
                           dtw_gamma=data.get('DTW gamma'),
                           transformer_layers=data.get('Transformer layers'),
                           transformer_heads=data.get('Transformer heads'),
                           window_width=data.get('Window size'),
                           window_stride=data.get('Actual stride pixels'), seed=data.get('Seed'))
        epoch_rows = workbook['EpochMetrics'].iter_rows(values_only=True)
        epoch_headers = next(epoch_rows, ())
        for row in epoch_rows:
            data = dict(zip(epoch_headers, row))
            trial = trials.get(data.get('Trial_ID'))
            if trial is None:
                continue
            epoch = int(data['Epoch'])
            while len(trial['history']) < epoch:
                trial['history'].append(dict(epoch=len(trial['history']) + 1))
            stats = {key: data.get(label) for key, label in zip(METRIC_KEYS, METRICS)}
            stats['gradients'] = dict(cnn_mean=stats.get('local_gradient_norm'),
                transformer_mean=stats.get('contextual_gradient_norm'),
                fusion_mean=stats.get('fusion_gradient_norm'))
            trial['history'][epoch - 1]['validation' if data['Split'] == 'val' else 'train'] = stats
        for trial in trials.values():
            if trial['history']:
                trial['history'][-1]['model_parameter_count'] = next(
                    (row[TRIAL_HEADERS.index('Model parameter count')]
                     for row in workbook['Trials'].iter_rows(min_row=2, values_only=True)
                     if row[0] == trial['id']), None)
    workbook.close()
    return trials


def save_trial(path, *, trial_id, status, config, history, checkpoint=None,
               log_path=None, error=None, search_space=None):
    """Upsert one trial; write every epoch, rank, refresh dashboard, and fsync."""
    path = Path(path)
    with _locked(path):
        trials = _read_existing(path)
        if trial_id is not None:
            trials[trial_id] = dict(id=trial_id, status=status, config=dict(config),
                                    history=list(history), checkpoint=checkpoint,
                                    log_path=log_path, error=error)
        ranked = sorted(trials.values(), key=cmp_to_key(_sort_trials))
        book = Workbook()
        ranked_sheet = book.active
        ranked_sheet.title = 'RankedResults'
        ranked_rows = []
        for rank, trial in enumerate(ranked, 1):
            ranked_rows.append(((rank if trial['status'] == 'COMPLETE' else None),) +
                               _ranked_row(trial))
        _style_table(ranked_sheet, RANKED_HEADERS, ranked_rows)
        trial_rows = []
        for trial in sorted(trials.values(), key=lambda item: item['id']):
            trial_rows.append(_config_values(trial['id'], trial['status'], trial['config']) +
                _summary(trial['history'], trial.get('checkpoint'), trial.get('log_path'),
                         trial['config'].get('seed')) +
                (json.dumps(trial['config'], sort_keys=True), trial.get('error'),
                 datetime.now(timezone.utc).isoformat()))
        _style_table(book.create_sheet('Trials'), TRIAL_HEADERS, trial_rows)
        epoch_rows = []
        for trial in sorted(trials.values(), key=lambda item: item['id']):
            for entry in trial['history']:
                for split in ('train', 'validation'):
                    if split in entry:
                        epoch_rows.append((trial['id'], entry['epoch'],
                                           'val' if split == 'validation' else 'train') +
                                          metric_values(entry[split], split))
        _style_table(book.create_sheet('EpochMetrics'), EPOCH_HEADERS, epoch_rows)
        _dashboard(book.create_sheet('Dashboard'), ranked)
        if search_space is None and path.exists():
            existing = load_workbook(path, read_only=True, data_only=True)
            if 'SearchSpace' in existing:
                search_space = {row[0]: json.loads(row[1]) for row in
                    existing['SearchSpace'].iter_rows(min_row=2, values_only=True) if row[0]}
            existing.close()
        space_rows = [(key, json.dumps(value), 'Optuna search choices')
                      for key, value in (search_space or {}).items()]
        space_sheet = book.create_sheet('SearchSpace')
        _style_table(space_sheet, ('Parameter', 'Values', 'Notes'), space_rows)
        space_sheet.column_dimensions['A'].width = 28
        space_sheet.column_dimensions['B'].width = 46
        space_sheet.column_dimensions['C'].width = 25
        fd, temporary = tempfile.mkstemp(prefix='.optuna_tracker_', suffix='.xlsx', dir=path.parent)
        os.close(fd)
        try:
            book.save(temporary)
            with open(temporary, 'rb') as stream:
                os.fsync(stream.fileno())
            os.replace(temporary, path)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)


def initialize(path=DEFAULT_TRACKER, search_space=None):
    """Create the five-sheet workbook before the first trial starts."""
    path = Path(path)
    if not path.exists():
        save_trial(path, trial_id=None, status='INITIALIZING', config={}, history=[],
                   search_space=search_space)
